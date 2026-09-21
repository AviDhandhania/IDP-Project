from pptx import Presentation
prs = Presentation("docs/presentations/Crypto-Agility_Navigator_Review2.pptx")
for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            texts.append(shape.text.replace("\n", " | ")[:100])
    print(f"--- Slide {i+1} ---")
    for t in texts:
        print(t)
